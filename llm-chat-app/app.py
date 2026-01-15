import os
import json
import uuid
import base64
import sys
import logging
import queue
import time
import threading
from datetime import datetime
from pathlib import Path
from flask import Flask, render_template, request, jsonify, Response, stream_with_context
from flask_cors import CORS


def sse_response(generator):
    """创建SSE响应，带有正确的headers防止缓冲"""
    return Response(
        stream_with_context(generator),
        mimetype='text/event-stream',
        headers={
            'Cache-Control': 'no-cache',
            'X-Accel-Buffering': 'no',
            'Connection': 'keep-alive'
        }
    )
from dotenv import load_dotenv
import requests
from werkzeug.utils import secure_filename

# 配置详细的日志输出
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

# 过滤掉频繁的监控请求日志
class MonitorLogFilter(logging.Filter):
    def filter(self, record):
        # 过滤掉包含 /monitor/ 的请求日志
        if hasattr(record, 'msg') and '/monitor/' in str(record.msg):
            return False
        return True

# 应用过滤器到 werkzeug 日志
werkzeug_logger = logging.getLogger('werkzeug')
werkzeug_logger.addFilter(MonitorLogFilter())

# 文件解析库
import PyPDF2
import docx
import openpyxl
import csv
import io

# 添加agent目录到路径
sys.path.insert(0, str(Path(__file__).parent.parent / 'agent-lightrag-app'))
from lightrag_agent import run_agent_stream

# 添加DSMC模块路径
sys.path.insert(0, str(Path(__file__).parent.parent / 'agent-dsmc'))

# Import SSE manager
from sse_manager import sse_manager

# Import config manager
from config_manager import get_config_manager

# 延迟导入DSMC模块（避免循环依赖）
dsmc_agent_instance = None
dsmc_detector_instance = None

def get_dsmc_agent():
    global dsmc_agent_instance
    if dsmc_agent_instance is None:
        from dsmc_agent import DSMCAgent
        dsmc_agent_instance = DSMCAgent()
    return dsmc_agent_instance

def get_dsmc_detector():
    global dsmc_detector_instance
    if dsmc_detector_instance is None:
        from keyword_detector import DSMCKeywordDetector
        dsmc_detector_instance = DSMCKeywordDetector()
    return dsmc_detector_instance

load_dotenv()

app = Flask(__name__)
CORS(app)

# 配置
API_URL = os.getenv('API_URL', 'https://api.mjdjourney.cn/v1')
API_KEY = os.getenv('API_KEY', '')
PORT = int(os.getenv('PORT', 21000))
MODELS = os.getenv('MODELS', 'claude-opus-4-5-20251001,gemini-3-pro-preview,deepseek-v3-250324').split(',')

# 数据存储路径
DATA_DIR = Path('data')
DATA_DIR.mkdir(exist_ok=True)
CONVERSATIONS_FILE = DATA_DIR / 'conversations.json'
UPLOADS_DIR = DATA_DIR / 'uploads'
UPLOADS_DIR.mkdir(exist_ok=True)
RAG_RESULTS_DIR = DATA_DIR / 'rag_results'
RAG_RESULTS_DIR.mkdir(exist_ok=True)

# DSMC会话存储
DSMC_SESSIONS_DIR = DATA_DIR / 'dsmc_sessions'
DSMC_SESSIONS_DIR.mkdir(exist_ok=True)

# Session locking for concurrent run prevention
session_locks = {}
session_lock = threading.Lock()


def acquire_session_lock(session_id):
    """Acquire lock for a session, returns True if successful, False if already locked"""
    with session_lock:
        if session_id in session_locks and session_locks[session_id]:
            return False
        session_locks[session_id] = True
        return True


def release_session_lock(session_id):
    """Release lock for a session"""
    with session_lock:
        session_locks[session_id] = False


def is_session_locked(session_id):
    """Check if session is locked"""
    with session_lock:
        return session_id in session_locks and session_locks[session_id]


# 允许的文件类型
ALLOWED_EXTENSIONS = {'txt', 'pdf', 'doc', 'docx', 'xls', 'xlsx', 'csv', 'md', 'json', 'py', 'js', 'html', 'css', 'xml', 'yaml', 'yml', 'sparta', 'in', 'dat', 'stl', 'obj', 'surf', 'grid'}


def load_conversations():
    """加载对话历史"""
    if CONVERSATIONS_FILE.exists():
        with open(CONVERSATIONS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def save_conversations(conversations):
    """保存对话历史"""
    with open(CONVERSATIONS_FILE, 'w', encoding='utf-8') as f:
        json.dump(conversations, f, ensure_ascii=False, indent=2)


def save_rag_result(conv_id, query, parsed_data, full_response):
    """保存RAG搜索结果"""
    # 为每个对话创建一个目录
    conv_rag_dir = RAG_RESULTS_DIR / conv_id
    conv_rag_dir.mkdir(exist_ok=True)
    
    # 生成文件名（使用时间戳）
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    result_file = conv_rag_dir / f"rag_{timestamp}.json"
    
    # 构建保存的数据
    result_data = {
        "query": query,
        "timestamp": datetime.now().isoformat(),
        "response": full_response,
        "rag_data": {
            "kg_entities": parsed_data.get('kg_entities', []) if parsed_data else [],
            "kg_relationships": parsed_data.get('kg_relationships', []) if parsed_data else [],
            "dc_content": parsed_data.get('dc_content', []) if parsed_data else [],
            "statistics": {
                "entities_count": len(parsed_data.get('kg_entities', [])) if parsed_data else 0,
                "relationships_count": len(parsed_data.get('kg_relationships', [])) if parsed_data else 0,
                "documents_count": len(parsed_data.get('dc_content', [])) if parsed_data else 0
            }
        }
    }
    
    # 保存到文件
    with open(result_file, 'w', encoding='utf-8') as f:
        json.dump(result_data, f, ensure_ascii=False, indent=2)
    
    # 同时更新一个索引文件，方便查看所有RAG搜索记录
    index_file = RAG_RESULTS_DIR / 'index.json'
    index_data = []
    if index_file.exists():
        with open(index_file, 'r', encoding='utf-8') as f:
            index_data = json.load(f)
    
    index_data.append({
        "conv_id": conv_id,
        "query": query[:100] + ('...' if len(query) > 100 else ''),
        "timestamp": datetime.now().isoformat(),
        "file": str(result_file.relative_to(RAG_RESULTS_DIR)),
        "statistics": result_data["rag_data"]["statistics"]
    })
    
    # 只保留最近1000条记录
    index_data = index_data[-1000:]
    
    with open(index_file, 'w', encoding='utf-8') as f:
        json.dump(index_data, f, ensure_ascii=False, indent=2)
    
    return str(result_file)


def extract_text_from_file(file_path, filename):
    """从各种文件格式中提取文本"""
    ext = filename.rsplit('.', 1)[-1].lower()
    
    try:
        if ext == 'txt' or ext in ['py', 'js', 'html', 'css', 'md', 'json', 'xml', 'yaml', 'yml', 'sparta', 'in', 'dat', 'stl', 'obj', 'surf', 'grid']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        
        elif ext == 'pdf':
            text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text.append(page.extract_text() or '')
            return '\n'.join(text)
        
        elif ext == 'docx':
            doc = docx.Document(file_path)
            return '\n'.join([para.text for para in doc.paragraphs])
        
        elif ext == 'xlsx' or ext == 'xls':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = []
            for sheet in wb.worksheets:
                text.append(f"=== Sheet: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_text = '\t'.join([str(cell) if cell else '' for cell in row])
                    text.append(row_text)
            return '\n'.join(text)
        
        elif ext == 'csv':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                return '\n'.join(['\t'.join(row) for row in reader])
        
        else:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
    
    except Exception as e:
        return f"[文件解析错误: {str(e)}]"


@app.route('/')
def index():
    return render_template('index.html', models=MODELS)


@app.route('/api/models')
def get_models():
    return jsonify(MODELS)


@app.route('/api/conversations')
def get_conversations():
    """获取所有对话列表"""
    conversations = load_conversations()
    # 返回对话列表，按时间倒序
    conv_list = []
    for conv_id, conv in conversations.items():
        conv_list.append({
            'id': conv_id,
            'title': conv.get('title', '新对话'),
            'created_at': conv.get('created_at', ''),
            'updated_at': conv.get('updated_at', ''),
            'model': conv.get('model', MODELS[0]),
            'rag_enabled': conv.get('rag_enabled', False)
        })
    conv_list.sort(key=lambda x: x.get('updated_at', ''), reverse=True)
    return jsonify(conv_list)


@app.route('/api/conversations', methods=['POST'])
def create_conversation():
    """创建新对话"""
    conversations = load_conversations()
    conv_id = str(uuid.uuid4())
    now = datetime.now().isoformat()
    model = request.json.get('model', MODELS[0])
    rag_enabled = request.json.get('rag_enabled', False)

    logger.info(f"📝 创建新对话: {conv_id}")
    logger.info(f"  模型: {model}")
    logger.info(f"  RAG开关: {'开启' if rag_enabled else '关闭'}")

    conversations[conv_id] = {
        'title': '新对话',
        'created_at': now,
        'updated_at': now,
        'model': model,
        'rag_enabled': rag_enabled,
        'messages': []
    }
    save_conversations(conversations)
    return jsonify({'id': conv_id, 'title': '新对话', 'created_at': now})


@app.route('/api/conversations/<conv_id>')
def get_conversation(conv_id):
    """获取单个对话详情"""
    conversations = load_conversations()
    if conv_id in conversations:
        conv = conversations[conv_id]
        conv['id'] = conv_id
        return jsonify(conv)
    return jsonify({'error': 'Conversation not found'}), 404


@app.route('/api/conversations/<conv_id>', methods=['DELETE'])
def delete_conversation(conv_id):
    """删除对话"""
    conversations = load_conversations()
    if conv_id in conversations:
        logger.info(f"🗑️  删除对话: {conv_id}")
        del conversations[conv_id]
        save_conversations(conversations)
        # 同时删除对应的RAG结果目录
        conv_rag_dir = RAG_RESULTS_DIR / conv_id
        if conv_rag_dir.exists():
            import shutil
            shutil.rmtree(conv_rag_dir)
            logger.info(f"  已删除RAG结果目录")
        return jsonify({'success': True})
    logger.warning(f"⚠️  对话不存在: {conv_id}")
    return jsonify({'error': 'Conversation not found'}), 404


@app.route('/api/conversations/<conv_id>/title', methods=['PUT'])
def update_conversation_title(conv_id):
    """更新对话标题"""
    conversations = load_conversations()
    if conv_id in conversations:
        conversations[conv_id]['title'] = request.json.get('title', '新对话')
        conversations[conv_id]['updated_at'] = datetime.now().isoformat()
        save_conversations(conversations)
        return jsonify({'success': True})
    return jsonify({'error': 'Conversation not found'}), 404


@app.route('/api/upload', methods=['POST'])
def upload_file():
    """上传文件并提取文本"""
    if 'file' not in request.files:
        logger.error("❌ 上传失败: 未提供文件")
        return jsonify({'error': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        logger.error("❌ 上传失败: 未选择文件")
        return jsonify({'error': 'No file selected'}), 400

    # 保留原始文件名用于显示和扩展名检测
    original_filename = file.filename
    # 获取文件扩展名
    ext = original_filename.rsplit('.', 1)[-1].lower() if '.' in original_filename else ''
    # 使用uuid生成安全的存储文件名，但保留原始扩展名
    safe_filename = f"{uuid.uuid4()}.{ext}" if ext else str(uuid.uuid4())

    logger.info(f"📤 收到文件上传: {original_filename} (扩展名: {ext})")

    file_path = UPLOADS_DIR / safe_filename
    file.save(file_path)
    logger.info(f"  保存到: {file_path}")

    # 提取文本（使用原始文件名来检测类型）
    logger.info(f"  正在提取文本...")
    text = extract_text_from_file(file_path, original_filename)
    logger.info(f"  提取完成: {len(text)} 字符")

    # 删除临时文件
    file_path.unlink()

    return jsonify({
        'filename': original_filename,
        'text': text[:50000]  # 限制文本长度
    })


# DSMC临时文件目录
DSMC_TEMP_DIR = DATA_DIR / 'dsmc_temp'
DSMC_TEMP_DIR.mkdir(exist_ok=True)

@app.route('/api/dsmc/upload', methods=['POST'])
def dsmc_upload_file():
    """
    处理DSMC相关文件上传
    根据文件类型决定处理方式：
    - LLM参考文件：提取文本内容返回
    - 工作目录文件：保存到临时目录，返回文件路径
    """
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    file_type = request.form.get('type', 'other')
    zone = request.form.get('zone', 'workspace')

    filename = secure_filename(file.filename)
    logger.info(f"📤 DSMC文件上传: {filename}, 类型: {file_type}, 区域: {zone}")

    if zone == 'llm':
        # LLM参考文件：提取文本内容
        temp_path = DSMC_TEMP_DIR / f"{uuid.uuid4()}_{filename}"
        file.save(temp_path)

        # 提取文本
        content = extract_text_from_file(temp_path, filename)

        # 删除临时文件
        temp_path.unlink()

        return jsonify({
            'success': True,
            'filename': filename,
            'type': file_type,
            'zone': 'llm',
            'content': content[:50000]  # 限制长度
        })
    else:
        # 工作目录文件：保存到临时目录
        unique_filename = f"{uuid.uuid4()}_{filename}"
        temp_path = DSMC_TEMP_DIR / unique_filename
        file.save(temp_path)

        logger.info(f"  保存到临时目录: {temp_path}")

        return jsonify({
            'success': True,
            'filename': filename,
            'type': file_type,
            'zone': 'workspace',
            'temp_path': str(temp_path)
        })


@app.route('/api/dsmc/validate-input', methods=['POST'])
def validate_input_file():
    """
    验证SPARTA输入文件
    - 语法检查
    - 依赖文件检查（.vss, .stl等）
    """
    data = request.json or {}
    content = data.get('content', '')
    filename = data.get('filename', 'input.sparta')

    if not content:
        return jsonify({'valid': False, 'error': '内容为空'}), 400

    logger.info(f"🔍 验证SPARTA输入文件: {filename}")

    try:
        agent = get_dsmc_agent()
        result = agent.validate_input_file(content)

        logger.info(f"  验证结果: {'通过' if result['valid'] else '失败'}")
        return jsonify(result)
    except Exception as e:
        logger.error(f"❌ 验证失败: {str(e)}", exc_info=True)
        return jsonify({'valid': False, 'error': str(e)}), 500


@app.route('/api/dsmc/upload-input', methods=['POST'])
def upload_sparta_input():
    """
    Upload and validate SPARTA input file

    Returns:
        {
            "valid": bool,
            "temp_id": str (if valid),
            "params": Dict (if valid),
            "preview": str (if valid),
            "stats": Dict (if valid),
            "errors": List[str] (if invalid),
            "warnings": List[str],
            "suggestions": List[str] (if invalid)
        }
    """
    try:
        # Check if file in request
        if 'file' not in request.files:
            return jsonify({"valid": False, "errors": ["No file uploaded"]}), 400

        file = request.files['file']

        if file.filename == '':
            return jsonify({"valid": False, "errors": ["Empty filename"]}), 400

        # Read file content
        try:
            content = file.read().decode('utf-8')
        except UnicodeDecodeError:
            return jsonify({"valid": False, "errors": ["File is not valid UTF-8 text"]}), 400

        # Validate using SpartaValidator
        validator = get_sparta_validator()
        validation_result = validator.validate(content)

        if validation_result['valid']:
            # Extract parameters
            params = validator.extract_parameters(content)

            # Generate temp ID
            temp_id = str(uuid.uuid4())

            # Save to temporary location
            temp_file = UPLOADS_DIR / f"{temp_id}.sparta"
            with open(temp_file, 'w', encoding='utf-8') as f:
                f.write(content)

            # Generate preview (first 500 chars)
            preview = content[:500]
            if len(content) > 500:
                preview += "\n..."

            # Count lines and commands
            lines = content.split('\n')
            commands = sum(1 for line in lines if line.strip() and not line.strip().startswith('#'))

            logger.info(f"✅ SPARTA文件上传成功: {file.filename}, temp_id: {temp_id}")

            return jsonify({
                "valid": True,
                "temp_id": temp_id,
                "params": params,
                "preview": preview,
                "stats": {
                    "lines": len(lines),
                    "commands": commands,
                    "filename": file.filename
                },
                "warnings": validation_result['warnings']
            })
        else:
            # Invalid file
            logger.warning(f"❌ SPARTA文件验证失败: {file.filename}")
            return jsonify({
                "valid": False,
                "errors": validation_result['errors'],
                "warnings": validation_result['warnings'],
                "suggestions": validation_result['suggestions']
            }), 400

    except Exception as e:
        logger.error(f"Upload failed: {e}", exc_info=True)
        return jsonify({"valid": False, "errors": [str(e)]}), 500


# Helper to get validator instance
sparta_validator_instance = None

def get_sparta_validator():
    """Get or create SpartaValidator instance"""
    global sparta_validator_instance
    if sparta_validator_instance is None:
        sys.path.insert(0, str(Path(__file__).parent.parent / 'agent-dsmc'))
        from sparta_validator import SpartaValidator
        sparta_validator_instance = SpartaValidator()
    return sparta_validator_instance


def generate_session_id() -> str:
    """Generate session ID: YYYYMMDD_HHMMSS_randomhex"""
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    random_hex = uuid.uuid4().hex[:8]
    return f"{timestamp}_{random_hex}"


def override_run_steps(content: str, max_steps: int) -> str:
    """
    Override run command in SPARTA input file

    Replaces: run 1000
    With: run <max_steps>
    """
    import re

    lines = []
    for line in content.split('\n'):
        # Match run command with any number of steps
        match = re.match(r'^(\s*run\s+)\d+(.*)$', line)
        if match:
            # Replace steps
            lines.append(f"{match.group(1)}{max_steps}{match.group(2)}")
        else:
            lines.append(line)

    return '\n'.join(lines)


@app.route('/api/dsmc/run-uploaded', methods=['POST'])
def run_uploaded_file_direct():
    """
    Run uploaded SPARTA file directly

    Request:
        {
            "temp_id": str,
            "num_cores": int,
            "max_steps": int,
            "max_memory_gb": int,
            "max_fix_attempts": int
        }

    Returns:
        {
            "success": bool,
            "session_id": str (if success),
            "message": str,
            "error": str (if failure)
        }
    """
    try:
        data = request.get_json()

        # Validate required fields
        required_fields = ['temp_id', 'num_cores', 'max_steps']
        for field in required_fields:
            if field not in data:
                return jsonify({"success": False, "error": f"Missing field: {field}"}), 400

        temp_id = data['temp_id']
        num_cores = int(data.get('num_cores', 4))
        max_steps = int(data.get('max_steps', 1000))
        max_memory_gb = int(data.get('max_memory_gb', 100))
        max_fix_attempts = int(data.get('max_fix_attempts', 3))

        # Read temp file
        temp_file = UPLOADS_DIR / f"{temp_id}.sparta"
        if not temp_file.exists():
            return jsonify({"success": False, "error": "Temporary file not found"}), 404

        with open(temp_file, 'r', encoding='utf-8') as f:
            content = f.read()

        # Override run command with max_steps
        modified_content = override_run_steps(content, max_steps)

        # Create session
        session_id = generate_session_id()

        # Create session directory
        session_dir = DSMC_SESSIONS_DIR / session_id
        session_dir.mkdir(exist_ok=True, parents=True)

        # Save input file
        input_file_path = session_dir / 'input.sparta'
        with open(input_file_path, 'w', encoding='utf-8') as f:
            f.write(modified_content)

        # Create metadata
        metadata = {
            "session_id": session_id,
            "created_at": datetime.now().isoformat(),
            "source": "uploaded",
            "original_filename": temp_file.name,
            "input_file": str(input_file_path),
            "status": "pending",
            "run_params": {
                "num_cores": num_cores,
                "max_steps": max_steps,
                "max_memory_gb": max_memory_gb,
                "max_fix_attempts": max_fix_attempts
            },
            "iterations": []
        }

        metadata_path = session_dir / 'metadata.json'
        with open(metadata_path, 'w', encoding='utf-8') as f:
            json.dump(metadata, f, indent=2)

        # Run simulation in background thread
        def run_async():
            try:
                logger.info(f"🚀 Starting uploaded file simulation: {session_id}")

                # Send simulation started event
                sse_manager.send_event(session_id, 'simulation_started', {
                    'session_id': session_id,
                    'max_steps': max_run_steps
                })

                # Update status
                metadata['status'] = 'running'
                with open(metadata_path, 'w') as f:
                    json.dump(metadata, f, indent=2)

                # Add agent-dsmc to path if not already there
                import sys
                agent_dsmc_path = str(Path(__file__).parent.parent / 'agent-dsmc')
                if agent_dsmc_path not in sys.path:
                    sys.path.insert(0, agent_dsmc_path)

                # Run simulation
                from sparta_runner import SPARTARunner
                runner = SPARTARunner()

                start_time = time.time()
                result = runner.run(
                    input_file=str(input_file_path),
                    session_id=session_id,
                    num_cores=num_cores,
                    max_memory_gb=max_memory_gb,
                    timeout=600
                )
                total_time = time.time() - start_time

                # Update status
                metadata['status'] = 'completed' if result.get('success', False) else 'failed'
                metadata['run_result'] = result
                metadata['completed_at'] = datetime.now().isoformat()

                with open(metadata_path, 'w') as f:
                    json.dump(metadata, f, indent=2)

                # Send completion event
                if result.get('success', False):
                    sse_manager.send_event(session_id, 'simulation_completed', {
                        'session_id': session_id,
                        'status': 'completed',
                        'total_time': total_time,
                        'result': result
                    })
                    logger.info(f"✅ Simulation completed: {session_id}, success=True")
                else:
                    sse_manager.send_event(session_id, 'simulation_failed', {
                        'session_id': session_id,
                        'status': 'failed',
                        'error': result.get('error', 'Unknown error')
                    })
                    logger.info(f"⚠️ Simulation completed with errors: {session_id}")

            except Exception as e:
                logger.error(f"❌ Simulation failed: {session_id}, error={e}", exc_info=True)

                metadata['status'] = 'failed'
                metadata['error'] = str(e)
                with open(metadata_path, 'w') as f:
                    json.dump(metadata, f, indent=2)

                # Send error event
                sse_manager.send_event(session_id, 'simulation_failed', {
                    'session_id': session_id,
                    'status': 'failed',
                    'error': str(e)
                })

        import threading
        thread = threading.Thread(target=run_async, daemon=True)
        thread.start()

        logger.info(f"✅ Session created: {session_id}, running in background")

        return jsonify({
            "success": True,
            "session_id": session_id,
            "message": "Simulation started"
        })

    except Exception as e:
        logger.error(f"Failed to run uploaded file: {e}", exc_info=True)
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/dsmc/run-uploaded-input', methods=['POST'])
def run_uploaded_input():
    """
    直接运行上传的SPARTA输入文件
    创建新会话并作为V1初始版本
    """
    data = request.json or {}
    conversation_id = data.get('conversation_id')
    input_content = data.get('input_content', '')
    filename = data.get('filename', 'input.sparta')

    if not input_content:
        return jsonify({'error': '输入内容为空'}), 400

    logger.info("=" * 80)
    logger.info("🚀 直接运行上传的输入文件")
    logger.info(f"  对话ID: {conversation_id}")
    logger.info(f"  文件名: {filename}")

    def generate():
        try:
            agent = get_dsmc_agent()

            # 创建新DSMC会话
            yield f"data: {json.dumps({'type': 'status', 'message': '📁 创建DSMC会话...'})}\n\n"

            session_result = agent.create_session_from_uploaded_file(
                conversation_id=conversation_id,
                input_content=input_content,
                source='uploaded'
            )

            session_id = session_result['session_id']

            yield f"data: {json.dumps({'type': 'session_created', 'session_id': session_id})}\n\n"
            yield f"data: {json.dumps({'type': 'iteration_created', 'iteration': session_result['iteration']})}\n\n"

            # 运行仿真
            yield f"data: {json.dumps({'type': 'status', 'message': '🚀 开始运行SPARTA仿真...'})}\n\n"

            for event in agent.run_simulation(session_id, auto_fix=True):
                yield f"data: {json.dumps(event)}\n\n"

            yield f"data: {json.dumps({'type': 'run_complete'})}\n\n"

        except Exception as e:
            logger.error(f"❌ 运行失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


@app.route('/api/rag/results')
def get_rag_results():
    """获取RAG搜索结果列表"""
    index_file = RAG_RESULTS_DIR / 'index.json'
    if index_file.exists():
        with open(index_file, 'r', encoding='utf-8') as f:
            return jsonify(json.load(f))
    return jsonify([])


@app.route('/api/rag/results/<conv_id>')
def get_conv_rag_results(conv_id):
    """获取特定对话的RAG搜索结果"""
    conv_rag_dir = RAG_RESULTS_DIR / conv_id
    if not conv_rag_dir.exists():
        return jsonify([])
    
    results = []
    for file in sorted(conv_rag_dir.glob('rag_*.json'), reverse=True):
        with open(file, 'r', encoding='utf-8') as f:
            data = json.load(f)
            data['filename'] = file.name
            results.append(data)
    
    return jsonify(results)


@app.route('/api/rag/result/<conv_id>/<filename>')
def get_rag_result_detail(conv_id, filename):
    """获取特定RAG搜索结果的详情"""
    result_file = RAG_RESULTS_DIR / conv_id / filename
    if result_file.exists():
        with open(result_file, 'r', encoding='utf-8') as f:
            return jsonify(json.load(f))
    return jsonify({'error': 'Result not found'}), 404


@app.route('/api/chat', methods=['POST'])
def chat():
    """发送消息并获取回复（流式）"""
    data = request.json
    conv_id = data.get('conversation_id')
    user_message = data.get('message', '')
    model = data.get('model', MODELS[0])
    images = data.get('images', [])  # base64编码的图片列表
    rag_enabled = data.get('rag_enabled', False)

    logger.info("=" * 80)
    logger.info(f"📨 收到聊天请求")
    logger.info(f"  会话ID: {conv_id}")
    logger.info(f"  消息长度: {len(user_message)} 字符")
    logger.info(f"  模型: {model}")
    logger.info(f"  RAG开关: {'开启' if rag_enabled else '关闭'}")
    logger.info(f"  图片数量: {len(images)}")
    logger.info(f"  消息内容: {user_message[:100]}{'...' if len(user_message) > 100 else ''}")

    conversations = load_conversations()

    if conv_id not in conversations:
        logger.error(f"❌ 会话不存在: {conv_id}")
        return jsonify({'error': 'Conversation not found'}), 404

    conv = conversations[conv_id]

    # 添加用户消息到历史
    conv['messages'].append({
        "role": "user",
        "content": user_message,
        "images": images,
        "timestamp": datetime.now().isoformat()
    })

    # 更新对话标题（如果是第一条消息）
    if len(conv['messages']) == 1:
        conv['title'] = user_message[:30] + ('...' if len(user_message) > 30 else '')

    conv['updated_at'] = datetime.now().isoformat()
    conv['model'] = model
    conv['rag_enabled'] = rag_enabled
    save_conversations(conversations)

    # DSMC检测（优先级最高，独立于RAG开关）
    logger.info("🔍 开始DSMC意图检测...")
    detector = get_dsmc_detector()
    dsmc_result = detector.detect(user_message)
    logger.info(f"  检测结果: {'DSMC相关' if dsmc_result['is_dsmc'] else '普通对话'}")
    logger.info(f"  意图: {dsmc_result.get('intent', 'unknown')}")
    logger.info(f"  置信度: {dsmc_result.get('confidence', 0):.2%}")

    if dsmc_result['is_dsmc'] and dsmc_result['confidence'] > 0.6:
        logger.info("✅ 进入DSMC处理流程")
        return chat_with_dsmc(conv_id, user_message, model, conv, dsmc_result)

    # RAG开关：仅控制普通LLM对话是否使用LightRAG检索
    if rag_enabled:
        logger.info("✅ 进入RAG处理流程")
        return chat_with_rag(conv_id, user_message, model, conv)

    logger.info("✅ 进入普通LLM对话流程")

    # 构建消息内容
    content = []
    if user_message:
        content.append({"type": "text", "text": user_message})

    for img in images:
        content.append({
            "type": "image_url",
            "image_url": {"url": img}
        })

    # 如果只有文本，简化消息格式
    if len(content) == 1 and content[0]["type"] == "text":
        user_msg = {"role": "user", "content": user_message}
    else:
        user_msg = {"role": "user", "content": content}

    # 构建API请求的消息历史
    api_messages = []
    for msg in conv['messages']:
        if msg['role'] == 'user':
            if msg.get('images'):
                content = [{"type": "text", "text": msg['content']}]
                for img in msg['images']:
                    content.append({"type": "image_url", "image_url": {"url": img}})
                api_messages.append({"role": "user", "content": content})
            else:
                api_messages.append({"role": "user", "content": msg['content']})
        else:
            api_messages.append({"role": "assistant", "content": msg['content']})

    logger.info(f"📤 调用LLM API: {API_URL}")
    logger.info(f"  历史消息数: {len(api_messages)}")

    def generate():
        try:
            response = requests.post(
                f"{API_URL}/chat/completions",
                headers={
                    "Authorization": f"Bearer {API_KEY}",
                    "Content-Type": "application/json"
                },
                json={
                    "model": model,
                    "messages": api_messages,
                    "stream": True
                },
                stream=True,
                timeout=120
            )

            full_response = ""
            chunk_count = 0
            for line in response.iter_lines():
                if line:
                    line = line.decode('utf-8')
                    if line.startswith('data: '):
                        data = line[6:]
                        if data == '[DONE]':
                            break
                        try:
                            chunk = json.loads(data)
                            if 'choices' in chunk and len(chunk['choices']) > 0:
                                delta = chunk['choices'][0].get('delta', {})
                                if 'content' in delta:
                                    content = delta['content']
                                    full_response += content
                                    chunk_count += 1
                                    yield f"data: {json.dumps({'type': 'content', 'content': content})}\n\n"
                        except json.JSONDecodeError:
                            continue

            logger.info(f"✅ LLM响应完成: {len(full_response)} 字符, {chunk_count} 个片段")

            # 保存助手回复
            conversations = load_conversations()
            if conv_id in conversations:
                conversations[conv_id]['messages'].append({
                    "role": "assistant",
                    "content": full_response,
                    "timestamp": datetime.now().isoformat()
                })
                save_conversations(conversations)

            yield f"data: {json.dumps({'type': 'done', 'done': True})}\n\n"

        except Exception as e:
            logger.error(f"❌ LLM调用失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


def chat_with_rag(conv_id, user_message, model, conv):
    """使用RAG Agent处理聊天请求"""

    logger.info("🔎 进入RAG处理流程")
    logger.info(f"  查询: {user_message[:100]}")
    logger.info(f"  模型: {model}")

    def generate():
        try:
            full_response = ""
            parsed_data = None

            # 调用agent的流式接口
            for event in run_agent_stream(
                query=user_message,
                mode="mix",
                model=model,
                api_url=API_URL,
                api_key=API_KEY
            ):
                event_type = event.get('type')

                if event_type == 'status':
                    # 发送状态更新
                    logger.info(f"  状态: {event.get('message', '')}")
                    yield f"data: {json.dumps(event)}\n\n"

                elif event_type == 'content':
                    # 发送内容片段
                    full_response += event.get('content', '')
                    yield f"data: {json.dumps(event)}\n\n"

                elif event_type == 'done':
                    # 保存完整响应
                    full_response = event.get('full_response', full_response)
                    parsed_data = event.get('parsed_data')
                    logger.info(f"✅ RAG处理完成")
                    if parsed_data:
                        logger.info(f"  知识图谱实体: {len(parsed_data.get('kg_entities', []))}")
                        logger.info(f"  知识图谱关系: {len(parsed_data.get('kg_relationships', []))}")
                        logger.info(f"  文档片段: {len(parsed_data.get('dc_content', []))}")

                elif event_type == 'error':
                    logger.error(f"❌ RAG处理失败: {event.get('error', '')}")
                    yield f"data: {json.dumps(event)}\n\n"
                    return

            # 保存RAG搜索结果到文件
            rag_result_file = None
            if parsed_data:
                rag_result_file = save_rag_result(conv_id, user_message, parsed_data, full_response)
                logger.info(f"  结果已保存: {rag_result_file}")

            # 保存助手回复
            conversations = load_conversations()
            if conv_id in conversations:
                conversations[conv_id]['messages'].append({
                    "role": "assistant",
                    "content": full_response,
                    "rag_data": {
                        "entities": len(parsed_data['kg_entities']) if parsed_data else 0,
                        "relationships": len(parsed_data['kg_relationships']) if parsed_data else 0,
                        "documents": len(parsed_data['dc_content']) if parsed_data else 0,
                        "result_file": rag_result_file
                    } if parsed_data else None,
                    "timestamp": datetime.now().isoformat()
                })
                save_conversations(conversations)

            yield f"data: {json.dumps({'type': 'done', 'done': True})}\n\n"

        except Exception as e:
            logger.error(f"❌ RAG处理异常: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


# ==================== DSMC相关函数和端点 ====================

def chat_with_dsmc(conv_id, user_message, model, conv, dsmc_result):
    """使用DSMC代理处理消息（流式）"""

    logger.info("🚀 进入DSMC处理流程")
    logger.info(f"  意图: {dsmc_result.get('intent', 'unknown')}")
    logger.info(f"  置信度: {dsmc_result.get('confidence', 0):.2%}")

    def generate():
        try:
            # 获取DSMC代理
            agent = get_dsmc_agent()

            # 处理DSMC查询（流式）
            full_response = ""
            parsed_data = None
            event_count = 0

            for event in agent.handle_dsmc_query(user_message, conv.get('messages', [])):
                event_type = event.get('type')
                event_count += 1

                if event_type == 'status':
                    logger.info(f"  📊 状态: {event.get('message', '')}")
                    yield f"data: {json.dumps(event)}\n\n"

                elif event_type == 'content':
                    full_response += event.get('content', '')
                    yield f"data: {json.dumps(event)}\n\n"

                elif event_type == 'parameter_form':
                    logger.info("  📝 发送参数表单")
                    logger.info(f"     表单字段: {', '.join(event.get('form', {}).keys())}")
                    # 发送参数表单事件
                    yield f"data: {json.dumps(event)}\n\n"
                    # 暂时结束（等待用户填写表单）
                    yield f"data: {json.dumps({'type': 'done', 'done': True})}\n\n"
                    return

                elif event_type == 'done':
                    logger.info(f"  ✅ DSMC处理完成, 共{event_count}个事件")
                    yield f"data: {json.dumps(event)}\n\n"

            # 保存助手回复到对话历史
            conversations = load_conversations()
            if conv_id in conversations:
                conversations[conv_id]['messages'].append({
                    "role": "assistant",
                    "content": full_response,
                    "timestamp": datetime.now().isoformat(),
                    "dsmc_result": dsmc_result
                })
                save_conversations(conversations)
                logger.info(f"  💾 已保存到对话历史")

        except Exception as e:
            logger.error(f"❌ DSMC处理异常: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


@app.route('/api/dsmc/detect', methods=['POST'])
def dsmc_detect():
    """检测消息是否为DSMC相关"""
    data = request.json
    message = data.get('message', '')

    detector = get_dsmc_detector()
    result = detector.detect(message)

    return jsonify(result)


@app.route('/api/dsmc/generate', methods=['POST'])
def dsmc_generate_input():
    """生成SPARTA输入文件（流式）"""
    data = request.json
    conv_id = data.get('conversation_id')
    parameters = data.get('parameters', {})
    llm_files = data.get('llm_files', [])
    workspace_files = data.get('workspace_files', [])

    logger.info("=" * 80)
    logger.info("⚙️  收到DSMC生成输入文件请求")
    logger.info(f"  会话ID: {conv_id}")
    logger.info(f"  参数:")
    for key, value in parameters.items():
        if isinstance(value, dict):
            logger.info(f"    {key}:")
            for k, v in value.items():
                logger.info(f"      {k}: {v}")
        else:
            logger.info(f"    {key}: {value}")
    if llm_files:
        logger.info(f"  LLM参考文件: {[f.get('filename') for f in llm_files]}")
    if workspace_files:
        logger.info(f"  工作目录文件: {[f.get('filename') for f in workspace_files]}")

    def generate():
        try:
            agent = get_dsmc_agent()
            event_count = 0
            done_event_sent = False

            for event in agent.generate_input_file(parameters, llm_files=llm_files, workspace_files=workspace_files):
                event_type = event.get('type')
                event_count += 1

                if event_type == 'status':
                    logger.info(f"  📊 {event.get('message', '')}")
                    event_json = json.dumps(event, ensure_ascii=False)
                    logger.debug(f"     发送事件: {event_json[:100]}")
                    yield f"data: {event_json}\n\n"

                elif event_type == 'done':
                    result = event.get('result', {})
                    logger.info(f"  ✅ 输入文件生成完成 (共{event_count}个事件)")
                    logger.info(f"     会话ID: {result.get('session_id')}")
                    logger.info(f"     输入文件长度: {len(result.get('input_file', ''))} 字符")
                    logger.info(f"     注释数量: {len(result.get('annotations', {}))}")
                    timing = result.get('timing', {})
                    if timing:
                        logger.info(f"     总耗时: {timing.get('total_time')}秒")
                        for step, time in timing.get('steps', {}).items():
                            logger.info(f"       - {step}: {time}秒")

                    # 保存DSMC输入文件到对话历史
                    if conv_id:
                        conversations = load_conversations()
                        if conv_id in conversations:
                            conversations[conv_id]['messages'].append({
                                "role": "assistant",
                                "content": "DSMC输入文件已生成",
                                "timestamp": datetime.now().isoformat(),
                                "dsmc_input_file": {
                                    "session_id": result.get('session_id'),
                                    "input_file": result.get('input_file'),
                                    "annotations": result.get('annotations'),
                                    "parameter_reasoning": result.get('parameter_reasoning'),
                                    "warnings": result.get('warnings'),
                                    "parameters": parameters
                                }
                            })
                            save_conversations(conversations)
                            logger.info(f"  💾 DSMC输入文件已保存到对话历史")

                    # 发送done事件
                    event_json = json.dumps(event, ensure_ascii=False)
                    logger.info(f"  📤 发送done事件 ({len(event_json)} 字节)")
                    logger.debug(f"     事件内容: {event_json[:200]}...")
                    yield f"data: {event_json}\n\n"
                    done_event_sent = True

                else:
                    # 其他类型事件
                    event_json = json.dumps(event, ensure_ascii=False)
                    yield f"data: {event_json}\n\n"

            if not done_event_sent:
                logger.warning("  ⚠️  未发送done事件，可能发生错误")
            else:
                logger.info("  ✅ 流式响应完成")

        except Exception as e:
            logger.error(f"❌ DSMC输入文件生成失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


@app.route('/api/dsmc/run', methods=['POST'])
def dsmc_run_simulation():
    """运行SPARTA仿真（流式）"""
    data = request.json
    session_id = data.get('session_id')
    conversation_id = data.get('conversation_id')  # 对话ID用于保存结果
    num_cores = data.get('num_cores', 4)
    max_steps = data.get('max_steps', 1000)
    max_memory_gb = data.get('max_memory_gb', None)  # 内存限制（GB）
    max_fix_attempts = data.get('max_fix_attempts', 3)  # 最大修复次数

    if not session_id:
        logger.error("❌ 缺少session_id参数")
        return jsonify({'error': 'session_id is required'}), 400

    logger.info("=" * 80)
    logger.info("🚀 收到DSMC运行仿真请求")
    logger.info(f"  会话ID: {session_id}")
    logger.info(f"  对话ID: {conversation_id}")
    logger.info(f"  CPU核数: {num_cores}")
    logger.info(f"  最大步数: {max_steps}")
    logger.info(f"  最大修复次数: {max_fix_attempts}")
    if max_memory_gb:
        logger.info(f"  内存限制: {max_memory_gb} GB")

    def generate():
        try:
            agent = get_dsmc_agent()
            agent.max_fix_attempts = max_fix_attempts  # 设置最大修复次数
            event_count = 0

            for event in agent.run_simulation(session_id, num_cores=num_cores, max_steps=max_steps, max_memory_gb=max_memory_gb):
                event_type = event.get('type')
                event_count += 1

                if event_type == 'status':
                    logger.info(f"  📊 {event.get('message', '')}")
                elif event_type == 'error':
                    logger.error(f"  ❌ {event.get('error', '')}")
                elif event_type == 'done':
                    result = event.get('result', {})
                    logger.info(f"  ✅ 仿真完成 (共{event_count}个事件)")
                    if result:
                        logger.info(f"     摘要: {result.get('summary', {})}")
                        logger.info(f"     图表数量: {len(result.get('plots', []))}")

                        # 保存仿真结果到对话消息
                        conversations = load_conversations()
                        if conversation_id and conversation_id in conversations:
                            conversations[conversation_id]['messages'].append({
                                "role": "assistant",
                                "content": "DSMC仿真完成",
                                "timestamp": datetime.now().isoformat(),
                                "dsmc_simulation_results": {
                                    "session_id": session_id,
                                    "summary": result.get('summary'),
                                    "plots": result.get('plots'),
                                    "interpretation": result.get('interpretation'),
                                    "suggestions": result.get('suggestions')
                                }
                            })
                            save_conversations(conversations)
                            logger.info(f"  💾 仿真结果已保存到对话 {conversation_id}")

                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"

        except Exception as e:
            logger.error(f"❌ DSMC仿真运行失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


@app.route('/api/dsmc/stop/<session_id>', methods=['POST'])
def dsmc_stop_simulation(session_id):
    """停止正在运行的SPARTA仿真"""
    logger.info("=" * 80)
    logger.info(f"⏹️ 收到停止仿真请求: {session_id}")

    try:
        agent = get_dsmc_agent()
        result = agent.stop_simulation(session_id)

        if result.get('success'):
            logger.info(f"  ✅ 仿真已停止")
        else:
            logger.warning(f"  ⚠️ {result.get('message')}")

        return jsonify(result)
    except Exception as e:
        logger.error(f"❌ 停止仿真失败: {str(e)}", exc_info=True)
        return jsonify({'success': False, 'message': str(e)}), 500


@app.route('/api/dsmc/status/<session_id>')
def dsmc_simulation_status(session_id):
    """检查仿真是否正在运行"""
    try:
        agent = get_dsmc_agent()
        is_running = agent.is_simulation_running(session_id)
        return jsonify({'session_id': session_id, 'is_running': is_running})
    except Exception as e:
        logger.error(f"❌ 检查仿真状态失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions')
def dsmc_list_sessions():
    """列出所有DSMC会话"""
    sessions = []

    if DSMC_SESSIONS_DIR.exists():
        for session_dir in DSMC_SESSIONS_DIR.iterdir():
            if session_dir.is_dir():
                metadata_file = session_dir / "metadata.json"
                if metadata_file.exists():
                    with open(metadata_file, 'r', encoding='utf-8') as f:
                        sessions.append(json.load(f))

    sessions.sort(key=lambda x: x.get('timestamp', ''), reverse=True)
    return jsonify(sessions)


@app.route('/api/dsmc/sessions/<session_id>')
def dsmc_get_session(session_id):
    """获取DSMC会话详情"""
    session_dir = DSMC_SESSIONS_DIR / session_id
    metadata_file = session_dir / "metadata.json"

    if not metadata_file.exists():
        return jsonify({'error': 'Session not found'}), 404

    with open(metadata_file, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 优先从实际的 input.sparta 文件读取最新内容
    # 这确保下载的是迭代修复后的最新版本
    input_file_path = session_dir / "input.sparta"
    if input_file_path.exists():
        with open(input_file_path, 'r', encoding='utf-8') as f:
            data['input_file'] = f.read()

    return jsonify(data)


@app.route('/api/dsmc/monitor/<session_id>')
def dsmc_monitor_session(session_id):
    """监控DSMC会话目录 - 获取文件列表和统计信息"""
    session_dir = DSMC_SESSIONS_DIR / session_id

    if not session_dir.exists():
        return jsonify({'error': 'Session not found', 'session_dir': str(session_dir)}), 404

    # 收集文件信息
    files = []
    type_stats = {}

    for file_path in session_dir.iterdir():
        if file_path.is_file():
            stat = file_path.stat()
            ext = file_path.suffix.lower() or '.unknown'
            file_info = {
                'name': file_path.name,
                'size': stat.st_size,
                'modified': datetime.fromtimestamp(stat.st_mtime).isoformat(),
                'type': ext
            }
            files.append(file_info)

            # 统计类型
            if ext not in type_stats:
                type_stats[ext] = {'count': 0, 'total_size': 0}
            type_stats[ext]['count'] += 1
            type_stats[ext]['total_size'] += stat.st_size

    # 按修改时间排序
    files.sort(key=lambda x: x['modified'], reverse=True)

    return jsonify({
        'session_id': session_id,
        'session_dir': str(session_dir),
        'total_files': len(files),
        'files': files,
        'type_stats': type_stats
    })


@app.route('/api/dsmc/monitor/<session_id>/log')
def dsmc_monitor_log(session_id):
    """监控DSMC日志文件 - 获取log.sparta内容（支持tail）"""
    session_dir = DSMC_SESSIONS_DIR / session_id
    log_file = session_dir / "log.sparta"

    # 也尝试查找其他可能的日志文件
    if not log_file.exists():
        for pattern in ['*.log', 'log.*']:
            logs = list(session_dir.glob(pattern))
            if logs:
                log_file = logs[0]
                break

    if not log_file.exists():
        return jsonify({
            'error': 'Log file not found',
            'session_dir': str(session_dir),
            'content': '',
            'line_count': 0,
            'file_size': 0
        })

    # 获取参数
    lines = request.args.get('lines', 100, type=int)  # 默认返回最后100行
    offset = request.args.get('offset', 0, type=int)  # 从末尾向前的偏移量

    try:
        stat = log_file.stat()
        file_size = stat.st_size

        with open(log_file, 'r', encoding='utf-8', errors='replace') as f:
            all_lines = f.readlines()
            total_lines = len(all_lines)

            # 计算返回的行范围
            start = max(0, total_lines - lines - offset)
            end = total_lines - offset if offset > 0 else total_lines

            content = ''.join(all_lines[start:end])

        return jsonify({
            'session_id': session_id,
            'log_file': log_file.name,
            'content': content,
            'line_count': total_lines,
            'file_size': file_size,
            'showing_lines': f"{start+1}-{end}",
            'modified': datetime.fromtimestamp(stat.st_mtime).isoformat()
        })
    except Exception as e:
        logger.error(f"读取日志文件失败: {e}")
        return jsonify({'error': str(e), 'content': '', 'line_count': 0})


@app.route('/api/dsmc/workdir', methods=['GET', 'PUT'])
def dsmc_workdir():
    """获取或设置DSMC工作目录"""
    if request.method == 'GET':
        return jsonify({
            'workdir': str(DSMC_SESSIONS_DIR),
            'exists': DSMC_SESSIONS_DIR.exists()
        })
    else:
        # PUT - 更新工作目录（仅用于显示，实际不会改变系统目录）
        data = request.json
        new_workdir = data.get('workdir', '')

        # 验证目录
        if new_workdir:
            new_path = Path(new_workdir)
            if not new_path.exists():
                return jsonify({'error': 'Directory does not exist'}), 400
            if not new_path.is_dir():
                return jsonify({'error': 'Path is not a directory'}), 400

        return jsonify({
            'workdir': new_workdir or str(DSMC_SESSIONS_DIR),
            'message': 'Work directory updated (display only)'
        })


@app.route('/api/dsmc/sessions/<session_id>/files/<path:filename>')
def dsmc_get_file_content(session_id, filename):
    """获取DSMC会话中特定文件的内容"""
    from flask import send_file
    session_dir = DSMC_SESSIONS_DIR / session_id
    file_path = session_dir / filename

    if not file_path.exists():
        return jsonify({'error': 'File not found'}), 404

    # 安全检查：确保文件在session目录内
    try:
        file_path.resolve().relative_to(session_dir.resolve())
    except ValueError:
        return jsonify({'error': 'Invalid file path'}), 403

    try:
        stat = file_path.stat()

        # 检查文件扩展名，确定是否是图片文件
        image_extensions = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp'}
        file_ext = file_path.suffix.lower()

        if file_ext in image_extensions:
            # 图片文件：直接以二进制返回，设置正确的MIME类型
            mime_types = {
                '.png': 'image/png',
                '.jpg': 'image/jpeg',
                '.jpeg': 'image/jpeg',
                '.gif': 'image/gif',
                '.bmp': 'image/bmp',
                '.webp': 'image/webp'
            }
            return send_file(
                file_path,
                mimetype=mime_types.get(file_ext, 'application/octet-stream')
            )

        # 非图片文件：限制大小并返回JSON格式
        if stat.st_size > 1024 * 1024:
            return jsonify({
                'error': 'File too large',
                'file_size': stat.st_size
            }), 413

        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()

        return jsonify({
            'filename': filename,
            'content': content,
            'file_size': stat.st_size,
            'modified': datetime.fromtimestamp(stat.st_mtime).isoformat()
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ==================== DSMC版本管理API ====================

@app.route('/api/dsmc/sessions/<session_id>/versions', methods=['GET'])
def get_session_versions(session_id):
    """获取会话的版本历史"""
    logger.info(f"📋 获取会话版本历史: {session_id}")

    try:
        agent = get_dsmc_agent()
        result = agent.get_version_history(session_id)

        if result.get("success"):
            logger.info(f"  ✅ 找到 {len(result.get('versions', []))} 个版本")
            return jsonify(result)
        else:
            logger.error(f"  ❌ 获取版本历史失败: {result.get('error')}")
            return jsonify(result), 404

    except Exception as e:
        logger.error(f"❌ 获取版本历史异常: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/versions/<int:version>', methods=['POST'])
def restore_session_version(session_id, version):
    """恢复到指定版本"""
    logger.info(f"🔄 恢复会话版本: {session_id} -> v{version}")

    try:
        agent = get_dsmc_agent()
        result = agent.restore_version(session_id, version)

        if result.get("success"):
            logger.info(f"  ✅ 已恢复到版本 v{version}")
            return jsonify(result)
        else:
            logger.error(f"  ❌ 恢复失败: {result.get('message')}")
            return jsonify(result), 400

    except Exception as e:
        logger.error(f"❌ 恢复版本异常: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/fix', methods=['POST'])
def manual_fix_session(session_id):
    """手动触发错误修复（流式）"""
    logger.info("=" * 80)
    logger.info(f"🔧 手动触发错误修复: {session_id}")

    def generate():
        try:
            agent = get_dsmc_agent()

            for event in agent.manual_fix(session_id):
                event_type = event.get('type')

                if event_type == 'status':
                    logger.info(f"  📊 {event.get('message', '')}")
                elif event_type == 'error':
                    logger.error(f"  ❌ {event.get('error', '')}")
                elif event_type == 'done':
                    result = event.get('result', {})
                    logger.info(f"  ✅ 修复完成，版本: v{result.get('version')}")

                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"

        except Exception as e:
            logger.error(f"❌ 手动修复异常: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


@app.route('/api/dsmc/sessions/<session_id>/run', methods=['POST'])
def dsmc_run_with_auto_fix(session_id):
    """运行SPARTA仿真（支持自动修复，流式）"""
    data = request.json or {}
    num_cores = data.get('num_cores', 4)
    max_steps = data.get('max_steps', 1000)
    auto_fix = data.get('auto_fix', True)

    logger.info("=" * 80)
    logger.info("🚀 收到DSMC运行仿真请求（支持自动修复）")
    logger.info(f"  会话ID: {session_id}")
    logger.info(f"  CPU核数: {num_cores}")
    logger.info(f"  最大步数: {max_steps}")
    logger.info(f"  自动修复: {auto_fix}")

    def generate():
        try:
            agent = get_dsmc_agent()
            event_count = 0

            for event in agent.run_simulation(session_id, num_cores=num_cores,
                                              max_steps=max_steps, auto_fix=auto_fix):
                event_type = event.get('type')
                event_count += 1

                if event_type == 'status':
                    logger.info(f"  📊 {event.get('message', '')}")
                elif event_type == 'error':
                    logger.error(f"  ❌ {event.get('error', '')}")
                elif event_type == 'fix_applied':
                    logger.info(f"  🔧 {event.get('message', '')}")
                elif event_type == 'done':
                    result = event.get('result', {})
                    logger.info(f"  ✅ 仿真完成 (共{event_count}个事件)")

                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"

        except Exception as e:
            logger.error(f"❌ DSMC仿真运行失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


# ==================== 迭代管理API ====================

@app.route('/api/dsmc/sessions/<session_id>/iterations', methods=['GET'])
def get_iterations(session_id):
    """获取会话的所有迭代记录"""
    logger.info(f"📋 获取会话迭代列表: {session_id}")

    try:
        # Check if session exists
        session_dir = DSMC_SESSIONS_DIR / session_id
        if not session_dir.exists():
            return jsonify({'success': False, 'error': 'Session not found'}), 404

        agent = get_dsmc_agent()
        iterations = agent.get_iterations(session_id)
        statistics = agent.get_session_statistics(session_id)

        # Get current iteration ID from metadata
        metadata_file = session_dir / 'metadata.json'
        current_iteration_id = None
        if metadata_file.exists():
            with open(metadata_file, 'r') as f:
                metadata = json.load(f)
                current_iteration_id = metadata.get('current_iteration_id')

        logger.info(f"  ✅ 找到 {len(iterations)} 个迭代")
        return jsonify({
            'success': True,
            'iterations': iterations,
            'current_iteration_id': current_iteration_id,
            'statistics': statistics
        })
    except Exception as e:
        logger.error(f"❌ 获取迭代列表失败: {str(e)}", exc_info=True)
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/iterations/<iteration_id>', methods=['GET'])
def get_single_iteration(session_id, iteration_id):
    """获取单个迭代的完整数据"""
    logger.info(f"📋 获取单个迭代: {session_id}/{iteration_id}")

    try:
        agent = get_dsmc_agent()
        iteration = agent.get_iteration(session_id, iteration_id)

        if iteration:
            logger.info(f"  ✅ 找到迭代，input_file长度: {len(iteration.get('input_file', ''))}")
            return jsonify(iteration)
        else:
            logger.warning(f"  ⚠️ 迭代不存在")
            return jsonify({'error': '迭代不存在'}), 404
    except Exception as e:
        logger.error(f"❌ 获取单个迭代失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/iterations', methods=['POST'])
def create_iteration(session_id):
    """创建新迭代（手动编辑方式）"""
    data = request.json or {}
    source = data.get('source', 'manual_edit')
    description = data.get('description', '手动编辑')
    input_file = data.get('input_file', '')
    overwrite_current = data.get('overwrite_current', False)  # 新增参数

    logger.info(f"📝 创建新迭代: {session_id}")
    logger.info(f"  来源: {source}")
    logger.info(f"  描述: {description[:50]}...")
    logger.info(f"  覆盖当前: {overwrite_current}")

    if not input_file:
        return jsonify({'error': '输入文件内容不能为空'}), 400

    try:
        agent = get_dsmc_agent()

        if overwrite_current:
            # 覆盖当前迭代
            iteration = agent.update_current_iteration(
                session_id,
                input_file,
                description
            )
        else:
            # 创建新迭代
            iteration = agent.create_iteration(
                session_id,
                input_file,
                source,
                description
            )

        if 'error' in iteration:
            return jsonify(iteration), 404

        # 同时更新input.sparta文件
        session_dir = DSMC_SESSIONS_DIR / session_id
        input_file_path = session_dir / "input.sparta"
        with open(input_file_path, 'w', encoding='utf-8') as f:
            f.write(input_file)

        logger.info(f"  ✅ 迭代操作成功: #{iteration.get('iteration_number')}")
        return jsonify(iteration)
    except Exception as e:
        logger.error(f"❌ 迭代操作失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/iterations/<iteration_id>', methods=['GET'])
def get_iteration(session_id, iteration_id):
    """获取单个迭代详情"""
    logger.info(f"📋 获取迭代详情: {session_id}/{iteration_id}")

    try:
        agent = get_dsmc_agent()
        iteration = agent.get_iteration(session_id, iteration_id)

        if not iteration:
            return jsonify({'error': 'Iteration not found'}), 404

        return jsonify(iteration)
    except Exception as e:
        logger.error(f"❌ 获取迭代详情失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/iterations/<iteration_id>', methods=['DELETE'])
def delete_iteration(session_id, iteration_id):
    """删除指定迭代"""
    delete_files = request.args.get('delete_files', 'false').lower() == 'true'

    logger.info(f"🗑️ 删除迭代: {session_id}/{iteration_id}")
    logger.info(f"  删除文件: {delete_files}")

    try:
        agent = get_dsmc_agent()
        result = agent.delete_iteration(session_id, iteration_id, delete_files)

        if 'error' in result:
            return jsonify(result), 404

        logger.info(f"  ✅ 删除成功")
        return jsonify(result)
    except Exception as e:
        logger.error(f"❌ 删除迭代失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/iterations/<iteration_id>/restore', methods=['POST'])
def restore_iteration(session_id, iteration_id):
    """Restore a previous iteration as the current active iteration"""
    try:
        session_dir = DSMC_SESSIONS_DIR / session_id
        if not session_dir.exists():
            return jsonify({"success": False, "error": "Session not found"}), 404

        # Load session metadata
        metadata_file = session_dir / 'metadata.json'
        if not metadata_file.exists():
            return jsonify({"success": False, "error": "Session metadata not found"}), 404

        with open(metadata_file, 'r') as f:
            metadata = json.load(f)

        # Check if current iteration is running
        current_iter_id = metadata.get('current_iteration_id')
        if current_iter_id:
            for iter_data in metadata.get('iterations', []):
                if iter_data.get('iteration_id') == current_iter_id:
                    if iter_data.get('status') in ['running', 'fixing']:
                        return jsonify({
                            "success": False,
                            "error": "Cannot restore while simulation is running. Please stop current iteration first."
                        }), 400

        # Find the iteration to restore
        target_iteration = None
        for iter_data in metadata.get('iterations', []):
            if iter_data.get('iteration_id') == iteration_id:
                target_iteration = iter_data
                break

        if not target_iteration:
            return jsonify({"success": False, "error": "Iteration not found"}), 404

        # Update current iteration ID
        metadata['current_iteration_id'] = iteration_id

        # Save updated metadata
        with open(metadata_file, 'w') as f:
            json.dump(metadata, f, indent=2)

        # Copy iteration input file to current if exists
        iteration_input = session_dir / f"{iteration_id}_input.sparta"
        if iteration_input.exists():
            current_input = session_dir / 'input.sparta'
            import shutil
            shutil.copy(iteration_input, current_input)

        logger.info(f"✅ Restored to iteration {iteration_id} (v{target_iteration.get('iteration_number', 'N/A')})")

        return jsonify({
            "success": True,
            "current_iteration_id": iteration_id,
            "message": f"Restored to iteration v{target_iteration.get('iteration_number', 'N/A')}"
        })

    except Exception as e:
        logger.error(f"Error restoring iteration {iteration_id}: {e}", exc_info=True)
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/events')
def session_events_stream(session_id):
    """Server-Sent Events endpoint for real-time session updates"""

    def event_stream():
        client_queue = queue.Queue(maxsize=50)
        sse_manager.add_client(session_id, client_queue)

        try:
            # Send initial connection event
            yield f"data: {json.dumps({'type': 'connected', 'session_id': session_id})}\n\n"

            # Send heartbeat every 30 seconds to keep connection alive
            last_heartbeat = time.time()

            while True:
                try:
                    # Check for new events (timeout to allow heartbeat)
                    message = client_queue.get(timeout=30)
                    yield f"data: {json.dumps(message)}\n\n"
                    last_heartbeat = time.time()

                except queue.Empty:
                    # Send heartbeat
                    if time.time() - last_heartbeat >= 30:
                        yield f"data: {json.dumps({'type': 'heartbeat'})}\n\n"
                        last_heartbeat = time.time()

        except GeneratorExit:
            # Client disconnected
            sse_manager.remove_client(session_id, client_queue)

    return Response(
        event_stream(),
        mimetype='text/event-stream',
        headers={
            'Cache-Control': 'no-cache',
            'X-Accel-Buffering': 'no',
            'Connection': 'keep-alive'
        }
    )


@app.route('/api/dsmc/sessions/<session_id>/statistics', methods=['GET'])
def get_session_statistics(session_id):
    """获取会话统计信息"""
    logger.info(f"📊 获取会话统计: {session_id}")

    try:
        agent = get_dsmc_agent()
        statistics = agent.get_session_statistics(session_id)
        return jsonify(statistics)
    except Exception as e:
        logger.error(f"❌ 获取统计失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/iterate', methods=['POST'])
def iterate_with_natural_language():
    """通过自然语言修改输入文件（流式）"""
    data = request.json or {}
    session_id = data.get('session_id')
    modification_request = data.get('modification_request', '')

    if not session_id:
        return jsonify({'error': '缺少session_id'}), 400
    if not modification_request:
        return jsonify({'error': '缺少修改需求'}), 400

    logger.info("=" * 80)
    logger.info("🤖 自然语言迭代请求")
    logger.info(f"  会话ID: {session_id}")
    logger.info(f"  修改需求: {modification_request[:100]}...")

    def generate():
        try:
            agent = get_dsmc_agent()

            for event in agent.iterate_with_natural_language(session_id, modification_request):
                event_type = event.get('type')

                if event_type == 'status':
                    logger.info(f"  📊 {event.get('message', '')}")
                elif event_type == 'error':
                    logger.error(f"  ❌ {event.get('error', '')}")
                elif event_type == 'done':
                    result = event.get('result', {})
                    iteration = result.get('iteration', {})
                    logger.info(f"  ✅ 迭代完成: #{iteration.get('iteration_number')}")

                    # 同时更新input.sparta文件
                    session_dir = DSMC_SESSIONS_DIR / session_id
                    input_file_path = session_dir / "input.sparta"
                    with open(input_file_path, 'w', encoding='utf-8') as f:
                        f.write(result.get('input_file', ''))

                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"

        except Exception as e:
            logger.error(f"❌ 自然语言迭代失败: {str(e)}", exc_info=True)
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return sse_response(generate())


# ==================== 下载功能API ====================

@app.route('/api/dsmc/sessions/<session_id>/download/input', methods=['GET'])
def download_input_file(session_id):
    """下载输入文件"""
    logger.info(f"📥 下载输入文件: {session_id}")

    session_dir = DSMC_SESSIONS_DIR / session_id
    input_file_path = session_dir / "input.sparta"

    if not input_file_path.exists():
        # 尝试从metadata获取
        metadata_file = session_dir / "metadata.json"
        if metadata_file.exists():
            with open(metadata_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                input_content = data.get('input_file', '')
                if input_content:
                    response = Response(
                        input_content,
                        mimetype='text/plain',
                        headers={
                            'Content-Disposition': f'attachment; filename=input_{session_id[:8]}.sparta'
                        }
                    )
                    return response
        return jsonify({'error': '输入文件不存在'}), 404

    with open(input_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    response = Response(
        content,
        mimetype='text/plain',
        headers={
            'Content-Disposition': f'attachment; filename=input_{session_id[:8]}.sparta'
        }
    )
    return response


@app.route('/api/dsmc/sessions/<session_id>/download/all', methods=['GET'])
def download_all_files(session_id):
    """打包下载所有文件（zip）"""
    import io
    from zipfile import ZipFile

    logger.info(f"📦 打包下载所有文件: {session_id}")

    session_dir = DSMC_SESSIONS_DIR / session_id

    if not session_dir.exists():
        return jsonify({'error': '会话不存在'}), 404

    # 创建内存中的zip文件
    zip_buffer = io.BytesIO()

    try:
        with ZipFile(zip_buffer, 'w') as zip_file:
            for file_path in session_dir.iterdir():
                if file_path.is_file():
                    zip_file.write(file_path, file_path.name)

        zip_buffer.seek(0)

        response = Response(
            zip_buffer.getvalue(),
            mimetype='application/zip',
            headers={
                'Content-Disposition': f'attachment; filename=dsmc_{session_id[:8]}.zip'
            }
        )
        logger.info(f"  ✅ 打包成功")
        return response

    except Exception as e:
        logger.error(f"❌ 打包失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/api/dsmc/sessions/<session_id>/files/<path:filename>/download', methods=['GET'])
def download_single_file(session_id, filename):
    """下载单个文件"""
    logger.info(f"📥 下载文件: {session_id}/{filename}")

    session_dir = DSMC_SESSIONS_DIR / session_id
    file_path = session_dir / filename

    if not file_path.exists():
        return jsonify({'error': '文件不存在'}), 404

    # 安全检查：确保文件在session目录内
    try:
        file_path.resolve().relative_to(session_dir.resolve())
    except ValueError:
        return jsonify({'error': '无效的文件路径'}), 403

    try:
        with open(file_path, 'rb') as f:
            content = f.read()

        # 根据文件类型设置MIME类型
        mime_type = 'application/octet-stream'
        ext = file_path.suffix.lower()
        if ext in ['.txt', '.sparta', '.log']:
            mime_type = 'text/plain'
        elif ext == '.json':
            mime_type = 'application/json'
        elif ext == '.png':
            mime_type = 'image/png'
        elif ext == '.dat':
            mime_type = 'application/octet-stream'

        response = Response(
            content,
            mimetype=mime_type,
            headers={
                'Content-Disposition': f'attachment; filename={filename}'
            }
        )
        return response

    except Exception as e:
        logger.error(f"❌ 下载文件失败: {str(e)}", exc_info=True)
        return jsonify({'error': str(e)}), 500


# ==================== Settings Management ====================

@app.route('/api/settings', methods=['GET'])
def get_settings():
    """Get current application settings"""
    try:
        config = get_config_manager()
        settings = config.get_all()

        # Mask sensitive keys
        masked_settings = {}
        sensitive_keys = ['API_KEY', 'TOKEN', 'SECRET', 'PASSWORD']

        for key, value in settings.items():
            if any(sensitive in key.upper() for sensitive in sensitive_keys):
                if len(value) > 8:
                    masked_settings[key] = value[:8] + '...'
                else:
                    masked_settings[key] = '***'
            else:
                masked_settings[key] = value

        return jsonify({
            "settings": masked_settings,
            "editable_keys": [
                'API_URL', 'LLM_MODEL', 'RAG_ENABLED', 'MAX_TOKENS',
                'DEFAULT_TEMPERATURE', 'DEFAULT_MAX_STEPS'
            ]
        })

    except Exception as e:
        logger.error(f"Error getting settings: {e}")
        return jsonify({"error": str(e)}), 500


@app.route('/api/settings', methods=['POST'])
def update_settings():
    """Update application settings (runtime or persistent)"""
    try:
        data = request.get_json()
        updates = data.get('updates', {})
        persist = data.get('persist', False)  # Whether to write to .env

        if not updates:
            return jsonify({"error": "No updates provided"}), 400

        config = get_config_manager()

        # Validate keys
        sensitive_keys = ['API_KEY', 'TOKEN', 'SECRET', 'PASSWORD']
        for key in updates.keys():
            if any(sensitive in key.upper() for sensitive in sensitive_keys):
                logger.warning(f"Updating sensitive key: {key}")

        if persist:
            # Update .env file permanently
            config.update_env_file(updates)
            message = "Settings updated and saved to .env"
        else:
            # Set runtime overrides only
            for key, value in updates.items():
                config.set_runtime(key, value)
            config.save_runtime_overrides()
            message = "Settings updated (runtime only, restart to revert)"

        return jsonify({
            "success": True,
            "message": message,
            "updated_keys": list(updates.keys()),
            "persisted": persist
        })

    except Exception as e:
        logger.error(f"Error updating settings: {e}")
        return jsonify({"error": str(e)}), 500


@app.route('/api/settings/test-connection', methods=['POST'])
def test_api_connection():
    """Test API connection with provided or current credentials"""
    try:
        data = request.get_json() or {}
        config = get_config_manager()

        # Use provided credentials or current config
        api_url = data.get('API_URL') or config.get('API_URL')
        api_key = data.get('API_KEY') or config.get('API_KEY')
        model = data.get('LLM_MODEL') or config.get('LLM_MODEL')

        # Test with a minimal API call
        import requests
        headers = {
            'x-api-key': api_key,
            'anthropic-version': '2023-06-01',
            'Content-Type': 'application/json'
        }

        # Simple test: send minimal message
        test_url = f"{api_url.rstrip('/')}/messages"
        payload = {
            'model': model,
            'max_tokens': 10,
            'messages': [{'role': 'user', 'content': 'test'}]
        }

        response = requests.post(test_url, headers=headers, json=payload, timeout=10)

        if response.status_code == 200:
            return jsonify({
                "success": True,
                "message": "API connection successful",
                "model": model
            })
        else:
            return jsonify({
                "success": False,
                "error": f"API returned {response.status_code}: {response.text[:200]}"
            }), 400

    except requests.exceptions.Timeout:
        return jsonify({
            "success": False,
            "error": "Connection timeout - check API_URL"
        }), 400
    except Exception as e:
        return jsonify({
            "success": False,
            "error": str(e)
        }), 400


@app.route('/api/log-client-error', methods=['POST'])
def log_client_error():
    """Log client-side errors for debugging"""
    try:
        data = request.get_json()
        logger.error(f"❌ Client error: {data.get('message')}")
        logger.error(f"   Stack: {data.get('stack', 'N/A')[:500]}")
        logger.error(f"   URL: {data.get('url')}")
        logger.error(f"   Timestamp: {data.get('timestamp')}")
        return jsonify({"success": True})
    except Exception as e:
        logger.error(f"Error logging client error: {e}")
        return jsonify({"error": str(e)}), 500


if __name__ == '__main__':
    logger.info("=" * 80)
    logger.info("🚀 启动LLM聊天应用服务器")
    logger.info("=" * 80)
    logger.info(f"📍 服务器地址: http://localhost:{PORT}")
    logger.info(f"🌐 监听地址: 0.0.0.0:{PORT}")
    logger.info(f"🤖 可用模型: {', '.join(MODELS)}")
    logger.info(f"🔗 API URL: {API_URL}")
    logger.info(f"📁 数据目录: {DATA_DIR}")
    logger.info(f"📊 RAG结果目录: {RAG_RESULTS_DIR}")
    logger.info(f"⚡ DSMC会话目录: {DSMC_SESSIONS_DIR}")
    logger.info(f"🐛 调试模式: 开启")
    logger.info("=" * 80)
    logger.info("✅ 服务器启动中...")
    logger.info("")

    app.run(host='0.0.0.0', port=PORT, debug=True)